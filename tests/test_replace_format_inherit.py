"""N15 回归（2026-09-12）：replace 后新文字必须**继承主 run 的字符格式**。

真实故障（用户实测）：对含多种格式的段落做文字修改后，**整段都变成新格式（默认字体/字号）**。
根因：`_apply_one` 的 replace 把原 runs 全部标删除，而新文字以**一个不带 rPr 的 `<w:ins>` run**
追加 → 字体/字号/加粗等字符格式全部丢失。

修复：新 ins run **深拷贝主 run（文本最长者）的 `<w:rPr>`**，替换后的段落保持其主体字体/字号。
（段内混合格式的逐 run 精细保留——按 diff 边界拆分 runs——列为后续增强，见待修清单第二十三节。）
"""

from __future__ import annotations

import zipfile

from docx import Document
from docx.shared import Pt

from papermate.services.docx_edit import apply_edits


def test_replace_inherits_main_run_format(tmp_path) -> None:  # noqa: ANN001
    """多 run 段落（主 run 14pt/Arial + 次 run 9pt）整体替换 → 新文字继承 14pt/Arial。"""
    p = tmp_path / "a.docx"
    d = Document()
    para = d.add_paragraph()
    main = para.add_run("这是一段很长的正文文本，使用十四号字体。" * 2)  # 主 run（文本最长）
    main.font.size = Pt(14)
    main.font.name = "Arial"
    minor = para.add_run("短。")  # 次要 run（9pt），验证取的是「主」run 而非第一个
    minor.font.size = Pt(9)
    d.save(str(p))

    out = apply_edits(p, [{"pid": "P0", "op": "replace", "text": "替换后的全新文本。", "quote": "这是一段很长的正文文本"}])
    assert out["applied"] == [{"pid": "P0", "op": "replace"}]

    with zipfile.ZipFile(p) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    start = xml.find("<w:ins")
    assert start != -1, "应存在 <w:ins> 修订元素"
    ins_xml = xml[start : xml.find("</w:ins>", start) + 8]
    assert 'w:val="28"' in ins_xml, "新文字应继承主 run 的字号（14pt = 28 半点）"
    assert "Arial" in ins_xml, "新文字应继承主 run 的字体"
    assert 'w:val="18"' not in ins_xml, "不应继承次要 run 的 9pt（18 半点）"
