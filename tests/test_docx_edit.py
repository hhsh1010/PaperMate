"""C2·P1 · 改稿直写 docx 测试

覆盖：
- 段落视图：P{} 线性编号、空段/域代码段不可编辑
- track-changes 写入：replace / append / insert_after，原文本以 delText 保留（可逆）
- quote 校验阀：宁拒不错；pid 越界 / 空 text / 非法 op / 域段 / 非文档 → rejected / 400
- round-trip 保真：apply 后 zip.testzip + Document 重开验证

数据层纯本地（内存生成 docx + 临时目录），不触发大模型 / 网络，可离线运行。
"""

from __future__ import annotations

import io
import os
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document

from papermate.services import docx_edit
from papermate.services.docx_edit import (
    DocxEditError,
    DocxNotEditable,
    _para_text,
    apply_edits,
    list_paragraphs,
)


def make_docx(texts: list[str], path) -> None:
    """现场生成测试 docx（每条文本一个段落，空串生成空段）。"""
    d = Document()
    for t in texts:
        d.add_paragraph(t)
    d.save(str(path))


def make_docx_bytes(texts: list[str]) -> bytes:
    buf = io.BytesIO()
    d = Document()
    for t in texts:
        d.add_paragraph(t)
    d.save(buf)
    return buf.getvalue()


def _full_text(path) -> str:
    """段落后全文（含 delText 内的旧文字）：round-trip 检查用。"""
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return xml


def _para_accepted_text(path, para_idx: int = 0) -> str:
    """「接受全部修订」后的段落文本：w:t 保留、w:delText 忽略。"""
    from docx.oxml.ns import qn

    p = Document(str(path)).paragraphs[para_idx]
    return "".join((n.text or "") for n in p._p.iter() if n.tag == qn("w:t"))


def _para_rejected_text(path, para_idx: int = 0) -> str:
    """「拒绝全部修订」后的段落文本：w:delText 还原旧文、w:ins 内文字忽略。"""
    from docx.oxml.ns import qn

    p = Document(str(path)).paragraphs[para_idx]
    return "".join((n.text or "") for n in p._p.iter() if n.tag == qn("w:delText"))


# ---------------- 段落视图 ----------------

def test_list_paragraphs_indexes_and_editability(tmp_path):
    p = tmp_path / "a.docx"
    make_docx(["第一段内容。", "", "第三段内容。"], p)

    view = list_paragraphs(p)
    assert [v["pid"] for v in view] == ["P0", "P1", "P2"]
    assert view[0]["editable"] is True
    assert view[1]["editable"] is False and view[1]["empty"] is True  # 空段不可编辑
    assert view[2]["text"] == "第三段内容。"


# ---------------- track-changes 写入 ----------------

def test_replace_keeps_old_text_as_deleted(tmp_path):
    """replace：旧文字以 delText 保留（Word 中显示删除线、可拒绝），新文字以 ins 写入。"""
    p = tmp_path / "a.docx"
    make_docx(["这是旧的引言内容。", "第二段不动。"], p)

    out = apply_edits(p, [{"pid": "P0", "op": "replace", "text": "这是全新引言。", "quote": "旧的引言"}])
    assert out["applied"] == [{"pid": "P0", "op": "replace"}]
    assert out["rejected"] == []

    xml = _full_text(p)
    # N15 增强：新文字按 diff 拆分为多个 run —— 改用「接受/拒绝修订」的**语义断言**，
    # 不再依赖「整段在单个文本节点里」的字面 in（拆分后自然不成立）。
    assert _para_accepted_text(p, 0) == "这是全新引言。"  # 接受修订 == 新文本
    rejected = _para_rejected_text(p, 0)
    # 拒绝修订：被替换/删除的旧文字片段还原（equal 部分本就保留为普通文本，不计入 delText）
    assert "旧的" in rejected and "内容" in rejected
    assert "全新" not in rejected
    assert "<w:ins" in xml and "<w:del " in xml
    # round-trip：重开不抛异常，段落数不变（replace 不增删段落）
    assert len(Document(str(p)).paragraphs) == 2


def test_append_and_insert_after(tmp_path):
    p = tmp_path / "a.docx"
    make_docx(["原段。"], p)

    out = apply_edits(
        p,
        [
            {"pid": "P0", "op": "append", "text": "追加的一句。"},
            {"pid": "P0", "op": "insert_after", "text": "插队的新段落。"},
        ],
    )
    assert len(out["applied"]) == 2 and out["rejected"] == []

    paras = Document(str(p)).paragraphs
    assert len(paras) == 2  # insert_after 新增一段
    # 用「有效文本」断言（原生 .text 不读 <w:ins> 内容）
    assert "追加的一句。" in _para_text(paras[0])  # append 落在原段尾
    assert _para_text(paras[1]) == "插队的新段落。"  # insert_after 落在下一新段


def test_quote_mismatch_rejected(tmp_path):
    """quote 校验阀：quote 不在段落原文中 → rejected，且不产生任何写入。"""
    p = tmp_path / "a.docx"
    make_docx(["真实原文。"], p)
    before = p.read_bytes()

    out = apply_edits(p, [{"pid": "P0", "op": "replace", "text": "x", "quote": "不存在的引用"}])
    assert out["applied"] == []
    assert out["rejected"][0]["reason"] == "quote 与段落原文不匹配（安全拒绝）"
    assert p.read_bytes() == before  # 全部被拒 → 文件字节不变


def test_invalid_ops_rejected(tmp_path):
    p = tmp_path / "a.docx"
    make_docx(["内容。", ""], p)

    out = apply_edits(
        p,
        [
            {"pid": "P9", "op": "replace", "text": "x"},  # pid 越界
            {"pid": "P0", "op": "revoke", "text": "x"},  # 非法 op（delete/purge 均已合法化）
            {"pid": "P0", "op": "replace", "text": "  "},  # 空 text
            {"pid": "P1", "op": "replace", "text": "x"},  # 空段不可编辑
        ],
    )
    assert out["applied"] == []
    reasons = [r["reason"] for r in out["rejected"]]
    assert any("无法定位" in r for r in reasons)
    assert any("不支持" in r for r in reasons)
    assert any("text 为空" in r for r in reasons)
    assert any("不可编辑" in r for r in reasons)


def test_non_docx_rejected(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("# 不是 docx", encoding="utf-8")
    with pytest.raises(DocxNotEditable):
        apply_edits(p, [{"pid": "P0", "op": "replace", "text": "x"}])


def test_mixed_apply_and_reject(tmp_path):
    """部分成功部分拒绝：applied 照常写入，rejected 不影响其余 op。"""
    p = tmp_path / "a.docx"
    make_docx(["第一段。", "第二段。"], p)

    out = apply_edits(
        p,
        [
            {"pid": "P1", "op": "append", "text": "补充说明。"},
            {"pid": "P0", "op": "replace", "text": "x", "quote": "错的"},
        ],
    )
    assert len(out["applied"]) == 1 and len(out["rejected"]) == 1
    assert "补充说明。" in _para_text(Document(str(p)).paragraphs[1])


# ---------------- C3：delete / 表格 / numPr / 备份撤销 ----------------

def _make_list_docx(path) -> None:
    """生成含一个编号列表段的 docx（手动注入 numPr）。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    d = Document()
    p = d.add_paragraph("列表项一。")
    pPr = p._p.get_or_add_pPr()
    numPr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numId = OxmlElement("w:numId")
    numId.set(qn("w:val"), "1")
    numPr.append(ilvl)
    numPr.append(numId)
    pPr.append(numPr)
    d.save(str(path))


def test_delete_marks_whole_paragraph(tmp_path):
    """delete：段内文字 delText + 段落标记 del；track-changes 不物理删段（重开段数不变）。"""
    p = tmp_path / "a.docx"
    make_docx(["要删的段落。", "保留段落。"], p)

    out = apply_edits(p, [{"pid": "P0", "op": "delete", "text": "任意", "quote": "要删"}])
    assert out["applied"] == [{"pid": "P0", "op": "delete"}]

    xml = _full_text(p)
    assert "要删的段落。" in xml  # delText 保留（可拒绝修订恢复）
    assert "<w:del " in xml
    # 段落标记被标记删除（pPr/rPr/del）
    assert "<w:del" in xml.split("w:pPr", 1)[-1] or "w:del" in xml

    paras = Document(str(p)).paragraphs
    assert len(paras) == 2  # 不物理删除
    assert _para_text(paras[0]) == ""  # 有效文本为空（接受修订后整段消失）
    assert _para_text(paras[1]) == "保留段落。"

    # 视图层面：该段变 empty、不可编辑
    v = {x["pid"]: x for x in list_paragraphs(p)}
    assert v["P0"]["empty"] is True and v["P0"]["editable"] is False


def test_delete_accepts_empty_text(tmp_path):
    """delete 必须放行空 text（工作台删除前端恒传 text=""）。

    回归守卫：`apply_edits` 曾对全部 op 统一要求 text 非空，导致工作台删除被拒「text 为空」。
    该缺陷被既有测试掩盖——它们传哑元 text="任意" 绕过校验，与前端真实载荷不一致。
    """
    p = tmp_path / "a.docx"
    make_docx(["要删的段落。", "保留段落。"], p)

    # 前端真实载荷：delete + text="" + quote=当前段落原文
    out = apply_edits(p, [{"pid": "P0", "op": "delete", "text": "", "quote": "要删的段落。"}])
    assert out["applied"] == [{"pid": "P0", "op": "delete"}]
    assert out["rejected"] == []

    paras = Document(str(p)).paragraphs
    assert len(paras) == 2  # 不物理删除
    assert _para_text(paras[0]) == ""
    assert _para_text(paras[1]) == "保留段落。"


def test_non_delete_ops_still_require_text(tmp_path):
    """反向守卫：放宽 delete 不得误伤其余 op——replace/append/insert_after 仍须非空 text。"""
    for op in ("replace", "append", "insert_after"):
        p = tmp_path / f"{op}.docx"
        make_docx(["原文段落。"], p)
        out = apply_edits(p, [{"pid": "P0", "op": op, "text": ""}])
        assert out["applied"] == [], f"{op} 空 text 竟被接受"
        assert out["rejected"] and out["rejected"][0]["reason"] == "text 为空"


# ---------------- 落盘占用重试（WinError 32） ----------------

def test_apply_edits_retries_transient_locked_file(tmp_path, monkeypatch):
    """os.replace 遇瞬时共享冲突（WinError 32，预览流/杀软/索引短暂占用）→ 自动重试后成功。"""
    p = tmp_path / "a.docx"
    make_docx(["段落一。"], p)
    calls = {"n": 0}
    real_replace = os.replace

    def flaky_replace(src, dst):
        # ★ O-06 之后：撤销栈元数据（.baks.json）也走原子替换 ⇒ 计数必须**只针对 docx 本体**，
        #   否则栈写入会先消耗掉两次"占用"⇒ 计数口径失真（本测试的意图是验证 docx 的重试）。
        if Path(dst) == p:
            calls["n"] += 1
            if calls["n"] <= 2:  # 前两次模拟被占用
                raise PermissionError(32, "拒绝访问。")
        return real_replace(src, dst)

    monkeypatch.setattr(docx_edit.os, "replace", flaky_replace)
    out = apply_edits(p, [{"pid": "P0", "op": "replace", "text": "新文字。", "quote": "段落一"}])
    assert out["applied"] == [{"pid": "P0", "op": "replace"}]
    assert calls["n"] == 3  # 失败 2 次后第 3 次成功
    assert _para_text(Document(str(p)).paragraphs[0]) == "新文字。"
    assert not tmp_path.joinpath("a.docx.tmp").exists()  # tmp 已被 replace 消费


def test_apply_edits_persistently_locked_clear_error(tmp_path, monkeypatch):
    """持续被占用（如 Word 一直开着文档）→ 重试耗尽，抛带明确指引的业务错误；tmp 清理、原文件无损。"""
    p = tmp_path / "a.docx"
    make_docx(["段落一。"], p)

    def always_locked(src, dst):
        raise PermissionError(32, "拒绝访问。")

    monkeypatch.setattr(docx_edit.os, "replace", always_locked)
    with pytest.raises(DocxEditError) as ei:
        apply_edits(p, [{"pid": "P0", "op": "replace", "text": "新文字。", "quote": "段落一"}])
    assert "占用" in str(ei.value) and "Word" in str(ei.value)
    assert not tmp_path.joinpath("a.docx.tmp").exists()  # 失败路径清理 tmp
    assert _para_text(Document(str(p)).paragraphs[0]) == "段落一。"  # 原文件未被半途破坏


def test_delete_covers_ins_and_hyperlink_runs(tmp_path):
    """delete 必须覆盖嵌套 run（w:ins / w:hyperlink）。

    回归守卫：旧实现 `findall(w:r)` 只取直接子级 run——AI 改写留下的 ins 文字、
    真实文档的超链接/引用 run 漏标删除，表现为「徽标已删除，文字还在」。
    """
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = tmp_path / "a.docx"
    d = Document()
    para = d.add_paragraph()
    ins = OxmlElement("w:ins")  # 场景A：先前改写留下的 w:ins 嵌套 run
    r1 = OxmlElement("w:r")
    t1 = OxmlElement("w:t")
    t1.text = "改写后的文字"
    r1.append(t1)
    ins.append(r1)
    para._p.append(ins)
    hlink = OxmlElement("w:hyperlink")  # 场景B：超链接嵌套 run（真实文档常见）
    hlink.set(qn("r:id"), "rId1")
    r2 = OxmlElement("w:r")
    t2 = OxmlElement("w:t")
    t2.text = "链接文字"
    r2.append(t2)
    hlink.append(r2)
    para._p.append(hlink)
    d.add_paragraph("保留段落。")
    d.save(str(p))

    out = apply_edits(p, [{"pid": "P0", "op": "delete", "text": "", "quote": "改写后的文字"}])
    assert out["applied"] == [{"pid": "P0", "op": "delete"}]

    paras = Document(str(p)).paragraphs
    assert _para_text(paras[0]) == ""  # 嵌套文字全部标记删除（视图层消失）
    assert _para_text(paras[1]) == "保留段落。"
    xml = _full_text(p)
    assert "改写后的文字" in xml and "链接文字" in xml  # delText 保留痕迹（可拒绝修订恢复）


def test_replace_covers_nested_runs(tmp_path):
    """replace 对含嵌套 run 的段落：旧文字（含 ins/hyperlink）全部标记删除，仅新文字可见。"""
    from docx.oxml import OxmlElement

    p = tmp_path / "a.docx"
    d = Document()
    para = d.add_paragraph()
    ins = OxmlElement("w:ins")
    r1 = OxmlElement("w:r")
    t1 = OxmlElement("w:t")
    t1.text = "旧改写文字"
    r1.append(t1)
    ins.append(r1)
    para._p.append(ins)
    d.add_paragraph("保留段落。")
    d.save(str(p))

    out = apply_edits(p, [{"pid": "P0", "op": "replace", "text": "全新文字", "quote": "旧改写文字"}])
    assert out["applied"] == [{"pid": "P0", "op": "replace"}]
    assert _para_text(Document(str(p)).paragraphs[0]) == "全新文字"


# ---------------- purge：彻底删除（物理移除空壳段） ----------------

def test_purge_removes_deleted_shell_paragraph(tmp_path):
    """delete（留壳）→ purge：物理移除空壳段，段数减少、文档流不再包含该 pid。"""
    p = tmp_path / "a.docx"
    make_docx(["要删的段落。", "保留段落。"], p)

    out = apply_edits(p, [{"pid": "P0", "op": "delete", "text": "", "quote": "要删的段落。"}])
    assert out["applied"] == [{"pid": "P0", "op": "delete"}]
    # 空壳段：有效文本已清空 + 段落标记被标记删除（deleted 标记供前端识别）
    stream = {x["pid"]: x for x in docx_edit.list_stream(p)}
    assert stream["B0"]["empty"] is True

    assert stream["B0"]["deleted"] is True  # 前端「彻底删除」按钮的依据
    out2 = apply_edits(p, [{"pid": "B0", "op": "purge", "text": ""}])
    assert out2["applied"] == [{"pid": "B0", "op": "purge"}]

    doc = Document(str(p))
    assert len(doc.paragraphs) == 1  # 物理移除生效
    assert _para_text(doc.paragraphs[0]) == "保留段落。"
    # 注意：purge 后 body 重排，剩余段落的 B 地址会变——按内容断言而非旧 pid
    paras_after = [x for x in docx_edit.list_stream(p) if x["kind"] == "para"]
    assert len(paras_after) == 1 and paras_after[0]["text"] == "保留段落。"


def test_purge_rejected_on_text_paragraph(tmp_path):
    """安全阀：带有效文本的段落拒绝 purge（宁拒不错，只允许清除标记删除后的空壳）。"""
    p = tmp_path / "a.docx"
    make_docx(["还有内容的段落。"], p)
    out = apply_edits(p, [{"pid": "P0", "op": "purge", "text": ""}])
    assert out["applied"] == []
    assert "已清空" in out["rejected"][0]["reason"]
    assert len(Document(str(p)).paragraphs) == 1  # 原文件未动


def test_purge_rejected_in_table_cell(tmp_path):
    """安全阀：表格内段落拒绝 purge（物理移除会破坏单元格结构）。"""
    p = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("正文段。")
    table = d.add_table(rows=1, cols=1)
    table.rows[0].cells[0].paragraphs[0].add_run("单元格文字。")
    d.save(str(p))

    # 表格单元格段落在文档流视图的地址：B1.R0.C0.P0（正文段 B0，表格 B1）——cells 在表格条目内
    table_item = next(x for x in docx_edit.list_stream(p) if x["kind"] == "table")
    cell_pid = table_item["cells"][0]["pid"]
    assert cell_pid.endswith(".R0.C0.P0")
    # 先标记删除把单元格段清空成壳（绕过「已清空」阀），purge 才会命中表格阀——证明两道阀独立生效
    out_del = apply_edits(p, [{"pid": cell_pid, "op": "delete", "text": "", "quote": "单元格文字。"}])
    assert out_del["applied"] == [{"pid": cell_pid, "op": "delete"}]
    out = apply_edits(p, [{"pid": cell_pid, "op": "purge", "text": ""}])
    assert out["applied"] == []
    assert "表格内" in out["rejected"][0]["reason"]
    # 单元格结构未破坏：重开正常、段落仍在
    assert len(Document(str(p)).tables[0].rows[0].cells[0].paragraphs) == 1


def test_table_paragraph_addressing_and_edit(tmp_path):
    """表格段落：视图含 T 地址、replace 表格内段落生效、round-trip 通过。"""
    p = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("正文段。")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run("单元格甲。")
    table.rows[0].cells[1].paragraphs[0].add_run("单元格乙。")
    d.save(str(p))

    view = list_paragraphs(p)
    pids = [v["pid"] for v in view]
    assert "P0" in pids and "T0.R0.C0.P0" in pids and "T0.R0.C1.P0" in pids

    out = apply_edits(
        p,
        [{"pid": "T0.R0.C0.P0", "op": "replace", "text": "改写的单元格。", "quote": "单元格甲"}],
    )
    assert len(out["applied"]) == 1 and out["rejected"] == []

    doc = Document(str(p))
    cell_text = _para_text(doc.tables[0].rows[0].cells[0].paragraphs[0])
    assert "改写的单元格。" in cell_text
    # 旧文字在 XML 层以 delText 保留（可逆）；N15 拆分后可能分段 —— 用语义断言
    from docx.oxml.ns import qn  # noqa: PLC0415  与文件内其他用例一致的局部导入

    cell_para = doc.tables[0].rows[0].cells[0].paragraphs[0]
    rejected = "".join((n.text or "") for n in cell_para._p.iter() if n.tag == qn("w:delText"))
    # N15 拆分后 delText 片段取决于 diff 切分：只断言「有删除痕迹且片段来自原文」
    assert rejected.strip() != ""
    assert all(ch in "单元格甲。" for ch in rejected)


def test_insert_after_copies_numpr(tmp_path):
    """insert_after 在列表段后插新段：pPr（含 numPr）被复制，新段延续列表。"""
    from docx.oxml.ns import qn

    p = tmp_path / "a.docx"
    _make_list_docx(p)

    out = apply_edits(p, [{"pid": "P0", "op": "insert_after", "text": "列表项二。"}])
    assert out["applied"] == [{"pid": "P0", "op": "insert_after"}]

    paras = Document(str(p)).paragraphs
    assert len(paras) == 2
    new_ppr = paras[1]._p.find(qn("w:pPr"))
    assert new_ppr is not None and new_ppr.find(qn("w:numPr")) is not None
    assert _para_text(paras[1]) == "列表项二。"


def test_backup_created_and_restore(tmp_path):
    """备份/撤销（2026-09-06 新契约）：apply 自动建 .bak 并入 undo 栈；撤销多级回退；
    **撤销本身不再写入可撤销历史**——undo 栈空后再次撤销明确报错，而非把「撤销」回退掉（横跳）。"""
    p = tmp_path / "a.docx"
    make_docx(["原始内容。"], p)

    apply_edits(p, [{"pid": "P0", "op": "replace", "text": "改后一。", "quote": "原始"}])
    apply_edits(p, [{"pid": "P0", "op": "replace", "text": "改后二。", "quote": "改后一"}])
    baks = [b for b in p.parent.glob("a.docx.bak-*") if not b.name.endswith(".baks.json")]
    assert len(baks) == 2  # 两级 undo 快照

    # 撤销① → 回退最近一次操作（改后二 → 改后一）
    docx_edit.restore_from_backup(p)
    assert _para_text(Document(str(p)).paragraphs[0]) == "改后一。"

    # 撤销② → 逐级回退到原始
    docx_edit.restore_from_backup(p)
    assert _para_text(Document(str(p)).paragraphs[0]) == "原始内容。"

    # 撤销③ → undo 栈已空：明确报错，**状态保持原样**（旧实现此处会把「撤销」回退掉，来回横跳）
    with pytest.raises(docx_edit.DocxEditError):
        docx_edit.restore_from_backup(p)
    assert _para_text(Document(str(p)).paragraphs[0]) == "原始内容。"


def test_undo_then_new_op_clears_redo(tmp_path):
    """撤销后产生新操作：redo（被撤销的状态）随之失效——再次撤销回退的是新操作，而非横跳回旧状态。"""
    p = tmp_path / "a.docx"
    make_docx(["原始内容。"], p)

    apply_edits(p, [{"pid": "P0", "op": "replace", "text": "改后一。", "quote": "原始"}])
    docx_edit.restore_from_backup(p)  # 撤销 → 原始；改后一进入 redo 栈
    assert _para_text(Document(str(p)).paragraphs[0]) == "原始内容。"

    apply_edits(p, [{"pid": "P0", "op": "replace", "text": "改后二。", "quote": "原始"}])
    docx_edit.restore_from_backup(p)  # 撤销的是 op2 → 回到原始（若 redo 未清则会跳回「改后一」）
    assert _para_text(Document(str(p)).paragraphs[0]) == "原始内容。"

    with pytest.raises(docx_edit.DocxEditError):
        docx_edit.restore_from_backup(p)


def test_restore_adopts_legacy_backups(tmp_path):
    """兼容：栈机制启用前遗留的 .bak（无 .baks.json）→ 撤销时收编磁盘上最新一份作为栈底。"""
    import shutil as _shutil

    p = tmp_path / "a.docx"
    make_docx(["原始内容。"], p)
    # 手工造一份"旧格式"备份（不产生 .baks.json）
    _shutil.copy2(p, p.with_name(p.name + ".bak-20250101-000000-000000"))
    # 直接覆盖为改后内容（不走 apply_edits，避免生成 .baks.json，模拟旧版本遗留现场）
    make_docx(["改后。"], p)

    docx_edit.restore_from_backup(p)
    assert _para_text(Document(str(p)).paragraphs[0]) == "原始内容。"


def test_backup_retention_limit(tmp_path):
    """备份保留最近 _MAX_BACKUPS 份，最旧的自动清理。"""
    p = tmp_path / "a.docx"
    make_docx(["v0。"], p)
    for i in range(docx_edit._MAX_BACKUPS + 2):
        apply_edits(p, [{"pid": "P0", "op": "append", "text": f"追加{i}。"}])
    baks = list(p.parent.glob("a.docx.bak-*"))
    assert len(baks) == docx_edit._MAX_BACKUPS


def test_restore_without_backup_raises(tmp_path):
    p = tmp_path / "a.docx"
    make_docx(["内容。"], p)
    with pytest.raises(docx_edit.DocxEditError):
        docx_edit.restore_from_backup(p)


# ---------------- C4·P1：文档流编址（高保真渲染对齐） ----------------

def test_list_stream_addresses_in_document_order(tmp_path):
    """B 地址按 body 实际顺序（正文段与表格交错）；sectPr 不编址。"""
    p = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("第一段。")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run("表格甲。")
    table.rows[0].cells[1].paragraphs[0].add_run("表格乙。")
    d.add_paragraph("第二段。")
    d.save(str(p))

    stream = docx_edit.list_stream(p)
    kinds = [(it["kind"], it["pid"]) for it in stream]
    assert kinds == [("para", "B0"), ("table", "B1"), ("para", "B2")]
    assert stream[0]["text"] == "第一段。"
    cell_pids = [c["pid"] for c in stream[1]["cells"]]
    assert cell_pids == ["B1.R0.C0.P0", "B1.R0.C1.P0"]
    assert stream[1]["cells"][0]["text"] == "表格甲。"


def test_apply_edits_with_b_addresses(tmp_path):
    """B 地址可直接编辑：replace / insert_after 全走文档流地址。"""
    p = tmp_path / "a.docx"
    d = Document()
    d.add_paragraph("第一段。")
    d.add_paragraph("第二段。")
    d.save(str(p))

    out = apply_edits(
        p,
        [
            {"pid": "B0", "op": "replace", "text": "改写的第一段。", "quote": "第一段"},
            {"pid": "B1", "op": "insert_after", "text": "插在二段后。"},
        ],
    )
    assert len(out["applied"]) == 2 and out["rejected"] == []

    stream = docx_edit.list_stream(p)
    texts = [it.get("text", "") for it in stream if it["kind"] == "para"]
    assert texts[0] == "改写的第一段。"
    assert "插在二段后。" in texts[2]  # insert_after 新段进入文档流


def test_b_address_out_of_range(tmp_path):
    """越界 B 地址 rejected；sectPr 不占用编址索引。"""
    p = tmp_path / "a.docx"
    make_docx(["内容。"], p)

    out = apply_edits(p, [{"pid": "B99", "op": "replace", "text": "x"}])
    assert out["applied"] == [] and "无法定位" in out["rejected"][0]["reason"]

    # 文档带 sectPr（python-docx 保存必然含）：B0 仍指向第一段而非 sectPr
    stream = docx_edit.list_stream(p)
    assert len(stream) == 1 and stream[0]["kind"] == "para"


# ---------------- C4 增补：修改后保存导出（接受全部修订） ----------------

def test_accept_all_revisions(tmp_path):
    """导出干净版：ins 解包为普通 run、del（含 delText）消失、文档可正常重开。"""
    from docx.oxml.ns import qn

    p = tmp_path / "a.docx"
    make_docx(["要删的段落。", "保留段落。"], p)
    apply_edits(
        p,
        [
            {"pid": "P1", "op": "replace", "text": "改后的第二段。", "quote": "保留段落"},
            {"pid": "P0", "op": "delete", "text": "任意"},
        ],
    )
    # 原文件带修订痕迹
    assert "<w:ins" in _full_text(p) and "<w:del " in _full_text(p)

    data = docx_edit.accept_all_revisions(p)
    buf = io.BytesIO(data)
    assert zipfile.ZipFile(buf).testzip() is None  # zip 完整

    clean = Document(buf)
    # 无任何修订标记
    assert not clean.element.body.findall(".//" + qn("w:ins"))
    assert not clean.element.body.findall(".//" + qn("w:del"))
    texts = [_para_text(par) for par in clean.paragraphs]
    assert "改后的第二段。" in texts  # 新文字保留为普通文本
    assert "要删的段落。" not in texts  # 旧文字（delText）已随接受修订消失
    # 原文件不动（仍带痕迹）
    assert "<w:ins" in _full_text(p)


def test_accept_all_revisions_non_docx(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(DocxNotEditable):
        docx_edit.accept_all_revisions(p)


# ---------------- API 层 ----------------

@pytest.fixture
def client():
    import os
    import tempfile

    import papermate.config as cfg_mod

    tmp = tempfile.mkdtemp(prefix="papermate_docx_edit_test_")
    os.environ["PAPERMATE_DATA_DIR"] = tmp
    cfg_mod._settings = None
    from papermate.api.app import create_app

    app = create_app()
    from fastapi.testclient import TestClient

    with TestClient(app) as c:
        yield c


def _upload_docx(client, texts: list[str]) -> int:
    r = client.post(
        "/api/v1/documents",
        files={"file": ("paper.docx", make_docx_bytes(texts), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200, (r.status_code, r.text)
    return r.json()["id"]


def test_api_paragraphs_and_apply_edits(client):
    doc_id = _upload_docx(client, ["引言原文。", "方法原文。"])

    # 段落视图
    r = client.get(f"/api/v1/documents/{doc_id}/paragraphs")
    assert r.status_code == 200, r.text
    view = r.json()
    assert [v["pid"] for v in view] == ["P0", "P1"]
    assert view[0]["text"] == "引言原文。"

    # 应用编辑（含一条 quote 错误 → rejected）
    r = client.post(
        f"/api/v1/documents/{doc_id}/apply-edits",
        json={
            "edits": [
                {"pid": "P1", "op": "append", "text": "补充消融实验说明。"},
                {"pid": "P0", "op": "replace", "text": "x", "quote": "不匹配"},
            ]
        },
    )
    assert r.status_code == 200, r.text
    out = r.json()
    assert len(out["applied"]) == 1 and len(out["rejected"]) == 1
    assert out["paragraph_count_after"] == 2
    assert out["backup"]  # C3：有实际写入 → 自动备份
    # 修复后：同步 endpoint 走 BackgroundTasks 调度，reindex 必被触发（旧实现 create_task 静默失败）
    assert out["reindex_triggered"] is True

    # 再读段落视图：新文字已可见（w:t 文本），旧 quote 校验拒绝未误伤
    view2 = client.get(f"/api/v1/documents/{doc_id}/paragraphs").json()
    assert "补充消融实验说明。" in view2[1]["text"]
    assert "引言原文。" in view2[0]["text"]  # replace 被拒 → 原文完好


def test_api_restore_backup(client):
    """C3 API：apply（改后）→ restore（还原原文）→ 404 语义。"""
    doc_id = _upload_docx(client, ["原始内容。"])

    r = client.post(
        f"/api/v1/documents/{doc_id}/apply-edits",
        json={"edits": [{"pid": "P0", "op": "replace", "text": "改后内容。", "quote": "原始"}]},
    )
    assert r.status_code == 200, r.text

    r = client.post(f"/api/v1/documents/{doc_id}/restore-backup")
    assert r.status_code == 200, r.text
    assert r.json()["restored_from"]

    view = client.get(f"/api/v1/documents/{doc_id}/paragraphs").json()
    assert view[0]["text"] == "原始内容。"


def test_api_restore_backup_without_backup_400(client):
    doc_id = _upload_docx(client, ["内容。"])
    r = client.post(f"/api/v1/documents/{doc_id}/restore-backup")
    assert r.status_code == 400
    # 2026-09-06 新契约：无历史时提示"没有可撤销的操作"（旧文案"没有可用的备份文件"）
    assert "没有可撤销的操作" in r.json()["detail"]


def test_api_apply_edits_404_and_400(client):
    assert client.get("/api/v1/documents/999999/paragraphs").status_code == 404
    assert client.post("/api/v1/documents/999999/apply-edits", json={"edits": []}).status_code == 404
    assert client.post("/api/v1/documents/999999/restore-backup").status_code == 404

    doc_id = _upload_docx(client, ["内容。"])
    # 空 edits → 400
    r = client.post(f"/api/v1/documents/{doc_id}/apply-edits", json={"edits": []})
    assert r.status_code == 400


def test_api_apply_edits_triggers_reindex_background_task(client):
    """回归 P1：改稿 apply 后 RAG 重索引必须真正调度（曾因同步 endpoint 线程池内
    asyncio.create_task 抛 RuntimeError 被吞而静默失效）。

    ★ 2026-09-30 契约升级：改稿后调度的必须是 `reprocess`（完整重处理：重新解析→
    分块→嵌入），而不是 `reindex`——后者从 DB chunks 重建管线，chunks 未变 ⇒
    signature 秒回旧 state，检索到的是改稿前内容（探针
    `scripts/probe_edit_consistency.py` 证实；回归见 test_edit_reindex_consistency.py）。
    """
    doc_id = _upload_docx(client, ["原始内容。"])
    calls = {"n": 0}

    async def spy_reprocess(doc_id):  # noqa: ANN001
        calls["n"] += 1
        return None

    client.app.state.doc_svc.reprocess = spy_reprocess

    r = client.post(
        f"/api/v1/documents/{doc_id}/apply-edits",
        json={"edits": [{"pid": "P0", "op": "replace", "text": "改后内容。", "quote": "原始"}]},
    )
    assert r.status_code == 200, r.text
    assert r.json()["reindex_triggered"] is True
    # 关键断言：BackgroundTasks 在事件循环中被 await，reprocess 实际被调用
    assert calls["n"] >= 1


def test_api_restore_backup_triggers_reindex_background_task(client):
    """回归 P1：restore-backup 后同样必须触发完整重处理（契约同上，2026-09-30 升级）。"""
    doc_id = _upload_docx(client, ["原始内容。"])
    client.post(
        f"/api/v1/documents/{doc_id}/apply-edits",
        json={"edits": [{"pid": "P0", "op": "replace", "text": "改后内容。", "quote": "原始"}]},
    )
    calls = {"n": 0}

    async def spy_reprocess(doc_id):  # noqa: ANN001
        calls["n"] += 1
        return None

    client.app.state.doc_svc.reprocess = spy_reprocess

    r = client.post(f"/api/v1/documents/{doc_id}/restore-backup")
    assert r.status_code == 200, r.text
    assert r.json()["reindex_triggered"] is True
    assert calls["n"] >= 1


def test_prune_page_shots_prunes_oldest_when_over_budget():
    """回归 P1：page_shots 缓存清理必须真正删除最旧文件（曾因 async def 未 await 而永不执行）。"""
    from papermate.api.routes_documents import _prune_page_shots

    cache = Path(tempfile.mkdtemp())
    files = []
    for i in range(10):
        f = cache / f"p{i}.png"
        f.write_bytes(b"x")
        files.append(f)
    # 显式控制 mtime：p0 最旧 … p9 最新
    base = time.time() - 100
    for i, f in enumerate(files):
        os.utime(f, (base + i, base + i))

    # 总大小 10B > max_bytes=4 → 淘汰最旧 20% ≈ 2 个
    _prune_page_shots(cache, max_bytes=4)

    remaining = {p.name for p in cache.iterdir()}
    assert len(remaining) == 8, remaining
    assert "p0.png" not in remaining and "p1.png" not in remaining  # 最旧两个被删
    assert "p9.png" in remaining  # 最新保留


def test_api_export_tracked_and_clean(client):
    """C4 增补：导出带痕版（attachment + 修订标记）与干净版（接受全部修订）。"""

    def _doc_xml(data: bytes) -> str:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert z.testzip() is None
            return z.read("word/document.xml").decode("utf-8")

    doc_id = _upload_docx(client, ["原始内容。"])
    client.post(
        f"/api/v1/documents/{doc_id}/apply-edits",
        json={"edits": [{"pid": "P0", "op": "append", "text": "新增一句。"}]},
    )

    # 带痕版：200 + 含修订标记 + 附件下载
    r = client.get(f"/api/v1/documents/{doc_id}/download", params={"mode": "tracked"})
    assert r.status_code == 200
    assert "attachment" in r.headers["content-disposition"]
    xml = _doc_xml(r.content)
    assert "<w:ins" in xml and "新增一句。" in xml

    # 干净版：200 + 无修订标记 + zip 完整 + 新文字保留为普通文本
    r = client.get(f"/api/v1/documents/{doc_id}/download", params={"mode": "clean"})
    assert r.status_code == 200
    xml = _doc_xml(r.content)
    assert "<w:ins" not in xml and "<w:del " not in xml
    assert "新增一句。" in xml

    # 非法 mode / 不存在文档
    assert client.get(f"/api/v1/documents/{doc_id}/download", params={"mode": "x"}).status_code == 400
    assert client.get("/api/v1/documents/999999/download").status_code == 404


# ===================================================================== O-06（2026-10-01）
# 撤销栈数据安全：此前 `_load_baks_meta` `except: pass` 静默吞损坏（无日志、无痕迹），
# 后果：① restore 误报「没有可撤销的操作(400)」而磁盘上 `.bak-*` 仍真实存在；
# ② `_save_baks_meta`（非原子 write_text）随后**覆盖损坏栈** ⇒ 旧历史**永久失联**。


def _make_doc(tmp_path: Path, name: str = "a.docx") -> Path:
    """生成一个最小可用 docx（本文件既有方式：内存生成）。"""
    p = tmp_path / name
    buf = io.BytesIO()
    Document().save(buf)
    p.write_bytes(buf.getvalue())
    return p


def test_corrupt_baks_meta_is_preserved_and_warns(tmp_path, caplog):
    """损坏的撤销栈：显式告警 + 原文件改名保留（`.corrupt-*`）⇒ 不再被静默覆盖。"""
    import logging as _logging

    doc = _make_doc(tmp_path)
    meta_path = docx_edit._baks_meta_path(doc)
    meta_path.write_text("{bad-json", encoding="utf-8")

    with caplog.at_level(_logging.WARNING, logger="papermate.services.docx_edit"):
        meta = docx_edit._load_baks_meta(doc)

    assert meta == {"undo": [], "redo": []}
    assert any("撤销栈损坏" in r.message for r in caplog.records), "损坏必须显式告警（不再静默）"
    corrupts = list(tmp_path.glob("*.corrupt-*"))
    assert corrupts, "损坏栈必须改名保留（否则会被下次写入覆盖 ⇒ 历史永久失联）"
    assert corrupts[0].read_text(encoding="utf-8") == "{bad-json", "保留的是损坏现场本身"
    assert not meta_path.exists(), "改名后原路径不应残留"


def test_missing_baks_meta_is_a_normal_path(tmp_path, caplog):
    """文件**缺失**是正常路径（机制启用前的旧备份）—— 不打 warning（避免日志噪音）。"""
    import logging as _logging

    doc = _make_doc(tmp_path)
    with caplog.at_level(_logging.WARNING, logger="papermate.services.docx_edit"):
        meta = docx_edit._load_baks_meta(doc)
    assert meta == {"undo": [], "redo": []}
    assert not [r for r in caplog.records if r.levelno >= _logging.WARNING]


def test_restore_recovers_after_stack_corruption(tmp_path):
    """★★ 用户可见的核心收益：栈损坏后撤销**仍能工作**（修复前误报 400）。

    修复前：损坏栈 → 空栈 → 且 `meta_path.exists()` 为真 ⇒ **不**走"收编磁盘备份"兜底
    ⇒ 抛「没有可撤销的操作」，而磁盘上 `.bak-*` 明明存在。
    修复后：损坏栈被改名保留 ⇒ `meta_path` 不存在 ⇒ 收编兜底生效 ⇒ 正常还原。
    """
    doc = _make_doc(tmp_path)
    docx_edit._baks_meta_path(doc).write_text("{corrupt", encoding="utf-8")
    bak = docx_edit._bak_name(doc, "20261001-000000-000000")
    bak.write_bytes(b"old-content")

    out = docx_edit.restore_from_backup(doc)
    assert out["restored_from"] == bak.name
    assert doc.read_bytes() == b"old-content", "应从磁盘上真实存在的快照还原"


def test_save_baks_meta_keeps_original_on_failure(tmp_path, monkeypatch):
    """原子写：写入失败时**原文件保持完好**、tmp 不残留（修复前是截断式 write_text）。"""
    doc = _make_doc(tmp_path)
    meta_path = docx_edit._baks_meta_path(doc)
    meta_path.write_text('{"undo": ["keep"], "redo": []}', encoding="utf-8")

    def _boom(tmp, target):  # noqa: ANN001
        raise docx_edit.DocxEditError("模拟占用")

    monkeypatch.setattr(docx_edit, "_replace_with_retry", _boom)
    with pytest.raises(docx_edit.DocxEditError):
        docx_edit._save_baks_meta(doc, {"undo": [], "redo": []})

    assert meta_path.read_text(encoding="utf-8") == '{"undo": ["keep"], "redo": []}', (
        "写入失败后原栈文件被破坏 ⇒ 历史丢失（O-06 修复点）"
    )
    assert not list(tmp_path.glob("*.tmp")), "失败的 tmp 必须清理（不留垃圾）"


# --------------------------------------------------------------------------- #
# ★ 2026-10-10（用户实测：改稿后书签区内容跨边界挪动 + 插入文字继承上标）
#
# 场景复刻（期刊模板常见结构）：段落被书签边界切开，且以**上标** run 结尾——
#   [bmkStart 作者] 钱建国 / 1(上标) [bmkEnd 作者][bmkStart 单位] ，崔新程 / 2(上标)
# 旧实现的两个缺陷：
#   ① replace 的"原位重建"把整条生成序列插回**第一个 run 的位置** ⇒ 其后的 run
#      被拉到 bookmarkEnd 之前 ⇒ "单位"书签区被掏空（编辑点之外结构被改变）；
#   ② 段尾追加的 <w:ins> 继承**最后一个 run**（上标"2"）的格式 ⇒ 新增作者名
#      显示为上标小字（用户报告的"严重格式错误"）。
# --------------------------------------------------------------------------- #
def _make_bookmark_split_docx(path: Path) -> str:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph()
    pel = p._p

    def run(text: str, sup: bool = False) -> Any:
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:eastAsia"), "黑体")
        rpr.append(fonts)
        if sup:
            va = OxmlElement("w:vertAlign")
            va.set(qn("w:val"), "superscript")
            rpr.append(va)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        return r

    def bmark(kind: str, bid: int, name: str) -> Any:
        el = OxmlElement(f"w:bookmark{kind}")
        el.set(qn("w:id"), str(bid))
        el.set(qn("w:name"), name)
        return el

    # 段落标记格式（w:pPr/w:rPr）：黑体、无上标 —— 段尾追加应继承它
    ppr = OxmlElement("w:pPr")
    prpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:eastAsia"), "黑体")
    prpr.append(fonts)
    ppr.append(prpr)
    pel.insert(0, ppr)

    pel.append(bmark("Start", 1, "作者"))
    pel.append(run("钱建国"))
    pel.append(run("1", sup=True))
    pel.append(bmark("End", 1, "作者"))
    pel.append(bmark("Start", 2, "单位"))
    pel.append(run("，崔新程"))
    pel.append(run("2", sup=True))
    doc.save(path)
    return "钱建国1，崔新程2"


def test_replace_keeps_bookmark_boundaries_and_end_insert_inherits_para_mark(tmp_path):
    import docx
    from docx.oxml.ns import qn

    path = tmp_path / "b.docx"
    old = _make_bookmark_split_docx(path)
    pid = list_paragraphs(path)[0]["pid"]

    out = apply_edits(path, [{"pid": pid, "op": "replace", "text": old + " 费大夫", "quote": old}])
    assert out["applied"], out

    pel = docx.Document(str(path)).paragraphs[0]._p
    kids = list(pel)

    # ① 书签边界不动："，崔新程" 必须仍位于 [bmkEnd 作者] **之后**（不得跨边界左移）
    i_end = next(i for i, el in enumerate(kids) if el.tag == qn("w:bookmarkEnd"))
    i_cui = next(
        i
        for i, el in enumerate(kids)
        if el.tag == qn("w:r") and "崔新程" in (el.findtext(qn("w:t")) or "")
    )
    assert i_cui > i_end, "run 被挪过书签边界（P-08 回归！）"

    # ② 段尾追加的 ins 继承**段落标记**格式：无 vertAlign、字体为黑体
    ins = pel.find(qn("w:ins"))
    assert ins is not None, "应有 <w:ins> 修订"
    rpr = ins.find(f"{qn('w:r')}/{qn('w:rPr')}")
    assert rpr is not None
    assert rpr.find(qn("w:vertAlign")) is None, "段尾追加不得继承上标（P-08 回归！）"
    fonts = rpr.find(qn("w:rFonts"))
    assert fonts is not None and fonts.get(qn("w:eastAsia")) == "黑体"

    # ③ 文本正确（含前导空格的新增）
    assert _para_text(pel).endswith(" 费大夫")


def test_append_inherits_para_mark_format(tmp_path):
    import docx
    from docx.oxml.ns import qn

    path = tmp_path / "a.docx"
    _make_bookmark_split_docx(path)
    pid = list_paragraphs(path)[0]["pid"]

    out = apply_edits(path, [{"pid": pid, "op": "append", "text": " 费大夫"}])
    assert out["applied"], out

    pel = docx.Document(str(path)).paragraphs[0]._p
    ins = pel.find(qn("w:ins"))
    assert ins is not None, "append 应产生 <w:ins>"
    rpr = ins.find(f"{qn('w:r')}/{qn('w:rPr')}")
    assert rpr is not None, "append 必须带 rPr（继承段落标记格式，旧实现裸默认字体）"
    assert rpr.find(qn("w:vertAlign")) is None
    fonts = rpr.find(qn("w:rFonts"))
    assert fonts is not None and fonts.get(qn("w:eastAsia")) == "黑体"


# --------------------------------------------------------------------------- #
# ★ 2026-10-10（用户需求："只对文字进行修改，图和公式则不应进行修改"）
#
# 媒体段（含 w:drawing / w:object / w:pict 的段落）解锁编辑：
#   - replace 走"媒体锚定"：媒体 run 对文字 diff 贡献为零 ⇒ 旧实现会把它当
#     "无内容 run"移除 ⇒ 公式/图片丢失（这正是当初整段拒绝的原因）；
#   - 现在媒体 run **原位保留**，仅文字部分参与 diff 与重建。
# --------------------------------------------------------------------------- #
def _make_media_docx(path: Path) -> str:
    """图文混排段：[run"左侧为"][run(含 w:drawing)][run"，其中为高度。"]"""
    from docx.oxml import OxmlElement

    doc = Document()
    p = doc.add_paragraph()
    pel = p._p

    def trun(text: str) -> Any:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        return r

    pel.append(trun("左侧为"))
    media = OxmlElement("w:r")
    media.append(OxmlElement("w:drawing"))
    pel.append(media)
    pel.append(trun("，其中为高度。"))
    doc.save(path)
    return "左侧为，其中为高度。"


def test_replace_preserves_media_objects(tmp_path):
    import docx
    from docx.oxml.ns import qn

    path = tmp_path / "m.docx"
    old = _make_media_docx(path)
    pid = list_paragraphs(path)[0]["pid"]

    new = "左侧为输入特征图。其中为高度。"
    out = apply_edits(path, [{"pid": pid, "op": "replace", "text": new, "quote": old}])
    assert out["applied"], out

    pel = docx.Document(str(path)).paragraphs[0]._p
    # ① 媒体对象仍在（未被当作"无内容 run"移除——P-08 核心断言）
    assert len(pel.findall(f"{qn('w:r')}/{qn('w:drawing')}")) == 1, "媒体对象丢失（回归！）"
    # ② 文字已按 AI 改写更新
    assert _para_text(pel) == new
    # ③ 对象位置保持：媒体 run 仍在"前段文字"与"改写文字"之间（原位，未移动/未删除）
    kids = list(pel)
    i_media = next(i for i, el in enumerate(kids) if el.find(qn("w:drawing")) is not None)
    i_ins = next(i for i, el in enumerate(kids) if el.tag == qn("w:ins"))
    assert i_media < i_ins, "媒体对象应保持在插入文字之前（原位）"
    prev_text = "".join((t.text or "") for t in kids[i_media - 1].iter(qn("w:t")))
    assert prev_text == "左侧为", "媒体对象的前邻应仍是原文字"


def test_pure_media_paragraph_still_not_editable(tmp_path):
    """纯图段（无文字）仍不可编辑：没有文字可改，replace 必须被拒。"""
    from docx.oxml import OxmlElement

    path = tmp_path / "img.docx"
    doc = Document()
    p = doc.add_paragraph()
    media = OxmlElement("w:r")
    media.append(OxmlElement("w:drawing"))
    p._p.append(media)
    doc.save(path)

    items = list_paragraphs(path)
    assert items[0]["editable"] is False
    assert items[0]["empty"] is True

    out = apply_edits(path, [{"pid": items[0]["pid"], "op": "replace", "text": "x"}])
    assert out["applied"] == []
    assert "不可编辑" in out["rejected"][0]["reason"]
