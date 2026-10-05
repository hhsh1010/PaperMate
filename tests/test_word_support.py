"""Word（.docx/.doc）支持测试（2026-08-23）：解析 / 上传校验 / 文本预览 / 转换检测。

覆盖：
- DocxParser：段落 + Heading 样式标题 + 表格 → Section/raw_text；
- 上传：docx 端到端（生成内存 docx → 上传 → 解析文本预览）；.doc 老格式 → 400 提示转存；
- 魔数内容校验：假扩展名（非 zip 的 docx / 非 OLE 的 doc / 非 %PDF 的 pdf）→ 400；
- 扩展名白名单 / 大小上限；
- GET /documents/{id}/text 文本预览接口；
- conversion_service：引擎检测不抛异常、无引擎时 convert 抛 ConversionError。
"""

from __future__ import annotations

import io
import os
import tempfile

import pytest
from fastapi.testclient import TestClient

import papermate.config as cfg_mod
import papermate.runtime_settings as rs_mod
from lightcore.rag.parser import DocumentParser
from papermate.processing.parsers import DocParseError, DocxParser, ParserFactory
from papermate.services.conversion_service import ConversionError, convert, detect_engine


def _make_docx_bytes() -> bytes:
    """用 python-docx 在内存构造一个含标题/段落/表格的 .docx。"""
    from docx import Document

    doc = Document()
    doc.add_heading("Word 支持测试文档", level=1)
    doc.add_paragraph("这是第一段正文内容，用于验证 docx 解析与全文检索。")
    doc.add_heading("方法", level=2)
    doc.add_paragraph("我们提出一种用于文档解析的新方法。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "指标"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "准确率"
    table.cell(1, 1).text = "0.95"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------- DocxParser ----------------
def test_docx_parser_sections_and_raw(tmp_path):
    path = tmp_path / "t.docx"
    path.write_bytes(_make_docx_bytes())
    content = DocxParser().parse(path, title="t.docx")
    assert content.source_format == "docx"
    assert content.raw_text  # 全文非空
    assert "这是第一段正文" in content.raw_text
    headings = [s.heading for s in content.sections if s.heading]
    assert any("Word 支持测试文档" in (h or "") for h in headings)
    # 表格文本并入（可检索）
    assert any("准确率" in (s.text or "") for s in content.sections) or "准确率" in content.raw_text


def test_docx_parser_via_factory(tmp_path):
    path = tmp_path / "t.docx"
    path.write_bytes(_make_docx_bytes())
    parser = ParserFactory.for_file("t.docx")
    assert isinstance(parser, DocxParser)
    assert parser.parse(path).raw_text


def test_docx_parser_ole_doc_rejected(tmp_path):
    """真实 .doc（OLE 头 + 非 zip 内容）→ DocParseError 提示转存 .docx。"""
    path = tmp_path / "old.doc"
    path.write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 64)
    parser = ParserFactory.for_file("old.doc")
    with pytest.raises(DocParseError, match="转存为 .docx"):
        parser.parse(path)


def test_lightcore_document_parser_docx(tmp_path):
    path = tmp_path / "t.docx"
    path.write_bytes(_make_docx_bytes())
    blocks = DocumentParser().parse_docx(path)
    texts = " ".join(b.text for b in blocks)
    assert "这是第一段正文" in texts
    assert "准确率" in texts


# ---------------- 上传校验（扩展名 / 大小 / 魔数） ----------------
@pytest.fixture
def client():
    tmp = tempfile.mkdtemp(prefix="pm_word_")
    os.environ["PAPERMATE_DATA_DIR"] = tmp
    cfg_mod._settings = None
    rs_mod._state = None
    from papermate.api.app import create_app

    app = create_app()
    with TestClient(app) as c:
        yield c


def test_upload_docx_end_to_end(client):
    """docx 上传 → ready → 文本预览接口返回解析内容（在线预览数据源）。"""
    r = client.post(
        "/api/v1/documents",
        files={"file": ("报告.docx", _make_docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert r.status_code == 200, r.text
    doc = r.json()
    assert doc["status"] == "ready"
    t = client.get(f"/api/v1/documents/{doc['id']}/text")
    assert t.status_code == 200
    assert "这是第一段正文" in t.json()["text"]


def test_upload_old_doc_returns_400_with_hint(client):
    """真实 .doc（OLE 头）→ 解析失败 → 400 + 转存提示（不静默）。"""
    r = client.post(
        "/api/v1/documents",
        files={"file": ("旧文档.doc", bytes.fromhex("D0CF11E0A1B11AE1") + b"\x00" * 64, "application/msword")},
    )
    assert r.status_code == 400
    assert "docx" in r.json()["detail"]


def test_upload_bad_magic_rejected(client):
    """魔数内容校验：扩展名与内容不符 → 400。"""
    # 假 docx（内容不是 zip 头）
    r = client.post("/api/v1/documents", files={"file": ("假.docx", b"not a zip at all", "application/octet-stream")})
    assert r.status_code == 400
    assert "docx" in r.json()["detail"]
    # 假 pdf
    r = client.post("/api/v1/documents", files={"file": ("假.pdf", b"plain text not pdf", "application/pdf")})
    assert r.status_code == 400
    # 假 doc
    r = client.post("/api/v1/documents", files={"file": ("假.doc", b"not ole", "application/msword")})
    assert r.status_code == 400


def test_upload_unsupported_ext_and_size(client):
    # 不支持扩展名
    r = client.post("/api/v1/documents", files={"file": ("病毒.exe", b"MZ...", "application/octet-stream")})
    assert r.status_code == 400
    assert "不支持的文件类型" in r.json()["detail"]
    # 超 50MB：2026-09-06 起改为流式校验，语义正确的 413（Payload Too Large），
    # 且 Content-Length 预检在读取前即拦截（不再先全量入内存）。
    big = b"a" * (51 * 1024 * 1024)
    r = client.post("/api/v1/documents", files={"file": ("大.txt", big, "text/plain")})
    assert r.status_code == 413
    assert "50MB" in r.json()["detail"]


def test_precheck_validates_too(client):
    """precheck 同样走校验：非法类型 400。"""
    r = client.post("/api/v1/documents/precheck", files={"file": ("x.exe", b"MZ", "application/octet-stream")})
    assert r.status_code == 400


# ---------------- 转换检测（PDF ↔ Word 互转） ----------------
def test_detect_engine_never_raises():
    report = detect_engine()
    assert "libreoffice" in report and "word_com" in report
    assert isinstance(report["libreoffice"], bool)


def test_convert_without_engine_raises(tmp_path):
    """无引擎环境：convert 抛 ConversionError 且消息可读（如实说明依赖）。"""
    src = tmp_path / "a.docx"
    src.write_bytes(_make_docx_bytes())
    report = detect_engine()
    if report["libreoffice"] or report["word_com"]:
        pytest.skip("本机存在转换引擎，跳过无引擎断言")
    with pytest.raises(ConversionError):
        convert(src, "pdf", out_dir=tmp_path / "out")


def test_convert_invalid_pair(tmp_path):
    src = tmp_path / "a.pdf"
    src.write_bytes(b"%PDF-1.4")
    with pytest.raises(ConversionError, match="不支持的转换"):
        convert(src, "png", out_dir=tmp_path / "out")
