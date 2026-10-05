"""N15 增强（2026-09-12）：replace 的**逐 run 精确保留**——未变化部分格式与内容零改动。

背景：上一版修复（继承主 run 格式）解决了"整段退化为默认格式"，但段内混合格式
（如部分加粗/斜体）仍会丢失。本增强在**安全前提下**做逐 run 精确保留：
- 段内所有有效 run 都是段落直接子级 → 走 SequenceMatcher opcodes 逐块处理：
  equal 保留原样、replace/delete 以裁剪克隆包 <w:del>、新文字 <w:ins> 继承对应原 run 格式；
- 段内含嵌套结构（hyperlink / 既有 w:ins 等）→ **回退**整段替换（安全阀）。
"""

from __future__ import annotations

import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from papermate.services.docx_edit import apply_edits


def _xml_of(path) -> str:  # noqa: ANN001
    with zipfile.ZipFile(path) as z:
        return z.read("word/document.xml").decode("utf-8")


def test_replace_preserves_untouched_runs(tmp_path) -> None:  # noqa: ANN001
    """3 run 段落只改中间：首尾 run 原样保留（普通 w:t、字号不变），中间 del + ins。"""
    p = tmp_path / "a.docx"
    d = Document()
    para = d.add_paragraph()
    r1 = para.add_run("开头不变。")
    r1.font.size = Pt(14)
    r2 = para.add_run("中间要改的部分")
    r2.font.size = Pt(12)
    r2.bold = True
    r3 = para.add_run("结尾不变。")
    r3.font.size = Pt(14)
    d.save(str(p))

    out = apply_edits(
        p,
        [{"pid": "P0", "op": "replace", "text": "开头不变。中间已更新。结尾不变。", "quote": "中间要改的部分"}],
    )
    assert out["applied"] == [{"pid": "P0", "op": "replace"}]

    xml = _xml_of(p)
    # ⚠️ 必须从保存后的文件重读（apply_edits 是 round-trip：写盘后再验证），内存对象是旧内容
    reloaded = Document(str(p)).paragraphs[0]._p
    # ① 语义断言：接受修订 == 新文本；拒绝修订 == 只有变化部分被删
    accepted = "".join((n.text or "") for n in reloaded.iter() if n.tag == qn("w:t"))
    rejected = "".join((n.text or "") for n in reloaded.iter() if n.tag == qn("w:delText"))
    assert accepted == "开头不变。中间已更新。结尾不变。"
    assert rejected == "要改的部分"
    # ② 未变化部分：仍以普通 <w:t> 存在（未被标记删除 → 格式与内容零改动）
    assert ">开头不变。</w:t>" in xml  # 属性无关匹配（克隆带 xml:space）
    assert ">结尾不变。</w:t>" in xml
    # ③ 变化部分：旧文字成 delText，新文字在 <w:ins> 内
    assert ">要改的部分</w:delText>" in xml  # 属性无关匹配
    assert "已更新。" in xml  # ins 新文字（"中间"为保留克隆，语义断言已覆盖拼接结果）


def test_replace_with_nested_run_falls_back(tmp_path) -> None:  # noqa: ANN001
    """含 hyperlink（嵌套 run）的段落 → 安全回退：整段 del + 单 ins 继承主格式。"""
    p = tmp_path / "b.docx"
    d = Document()
    para = d.add_paragraph()
    r1 = para.add_run("正文文字，较长的一段。")
    r1.font.size = Pt(14)
    h = OxmlElement("w:hyperlink")
    hr = OxmlElement("w:r")
    ht = OxmlElement("w:t")
    ht.text = "链接文字"
    hr.append(ht)
    h.append(hr)
    para._p.append(h)  # noqa: SLF001  构造嵌套结构
    d.save(str(p))

    out = apply_edits(p, [{"pid": "P0", "op": "replace", "text": "整段替换后的文字。", "quote": "正文文字"}])
    assert out["applied"] == [{"pid": "P0", "op": "replace"}]

    xml = _xml_of(p)
    # 回退行为：全部旧文字（含链接文字）标删除，单条 <w:ins> 携带新文本
    assert "<w:delText>链接文字</w:delText>" in xml
    ins_start = xml.find("<w:ins")
    ins_xml = xml[ins_start : xml.find("</w:ins>", ins_start) + 8]
    assert "整段替换后的文字。" in ins_xml
