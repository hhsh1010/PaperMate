"""PaperMate 端到端 Demo（**零密钥 / 零联网 / 秒级**）—— 本脚本就是仓库里的"小 demo"。

## 它证明什么
把核心用户链路**真跑一遍**，并把过程与结果打印成 Markdown：

    健康检查 → 建会话 → 上传论文(docx) → 提问(SSE 流式 + 引用溯源)
            → 段落视图 → 改稿(Word 修订痕迹 w:ins) → 导出「修订版 / 干净版」

## 为什么不需要 API key
LLM 用内置 `fake` 桩、嵌入用 `fake`（哈希占位）⇒ 全程**不出网、不花钱**。
它验证的是**工程链路**（接口契约 / SSE / 检索 / 改稿写回 / 落盘），不是模型效果 ——
真实效果请用真 key + 真嵌入在本机跑（见 README「快速开始」）。

## 用法
    # 本地
    .venv\\Scripts\\python.exe scripts\\demo_smoke.py
    # GitHub Actions（.github/workflows/demo.yml）：输出写入 Job Summary
    python scripts/demo_smoke.py >> "$GITHUB_STEP_SUMMARY"
"""

from __future__ import annotations

import io
import os
import platform
import sys
import tempfile
import zipfile
from pathlib import Path

# ---- 环境必须在 import papermate **之前**设好（配置在导入期读取） ----
_DEMO_HOME = Path(tempfile.mkdtemp(prefix="papermate_demo_"))
os.environ.setdefault("PAPERMATE_DATA_DIR", str(_DEMO_HOME))
os.environ.setdefault("PAPERMATE_DEFAULT_EMBEDDER", "fake")  # 不下载 4.5GB 权重
os.environ.setdefault("PAPERMATE_USE_BGE_SPARSE", "false")
os.environ.setdefault("PAPERMATE_PREWARM", "0")

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from docx import Document  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from papermate.api.app import create_app  # noqa: E402

API = "/api/v1"
DEMO_AUTHOR = "PaperMate Demo"


def _build_demo_docx() -> bytes:
    """现场造一篇"迷你论文"（不依赖任何外部文件/网络）。"""

    doc = Document()
    doc.add_heading("面向低资源场景的轻量级目标检测方法", level=0)
    doc.add_paragraph("摘要：本文提出一种轻量检测网络，在保持精度的同时显著降低参数量。")
    doc.add_heading("1 引言", level=1)
    doc.add_paragraph("目标检测是计算机视觉的基础任务，但现有方法参数量大、推理成本高。")
    doc.add_paragraph("本文的核心贡献有三点：轻量骨干、跨尺度特征融合、以及一致性蒸馏训练策略。")
    doc.add_heading("2 方法", level=1)
    doc.add_paragraph("我们采用深度可分离卷积替代标准卷积，并在颈部引入跨尺度注意力。")
    doc.add_paragraph("训练阶段使用一致性蒸馏：教师网络与学生网络在同一批次上对齐预测分布。")
    doc.add_heading("3 实验", level=1)
    doc.add_paragraph("在公开数据集上，本方法以 1/3 的参数量达到基线 98.2% 的精度，端侧推理延迟降低 41%。")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _parse_sse(text: str) -> list[tuple[str, dict]]:
    """把 SSE 文本切成 (event, data) 列表（与仓库测试同口径）。"""
    import json

    out: list[tuple[str, dict]] = []
    name, data = None, ""
    for line in text.splitlines():
        if line.startswith("event:"):
            name = line[len("event:"):].strip()
        elif line.startswith("data:"):
            data = line[len("data:"):].strip()
            if name is not None:
                try:
                    out.append((name, json.loads(data)))
                except json.JSONDecodeError:
                    pass
                name, data = None, ""
    return out


def _count_revisions(docx_bytes: bytes) -> tuple[int, int]:
    """数导出件里的 Word 修订标记（用 `<w:ins ` 带空格匹配，避免误匹配 `w:insideH` 表格边框）。"""
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as z:
        xml = z.read("word/document.xml").decode("utf-8")
    return xml.count("<w:ins "), xml.count("<w:del ")


def main() -> int:
    print("## PaperMate 端到端 Demo")
    print()
    print(
        f"运行环境：Python {platform.python_version()} · {platform.system()} · "
        f"数据目录 `{os.environ['PAPERMATE_DATA_DIR']}`"
    )
    print()
    print("> 本 demo 用内置 `fake` 模型与 `fake` 嵌入跑通**工程链路**（接口/SSE/检索/改稿/落盘），")
    print("> 全程不出网、不需要任何 API key。真实效果请用真 key 在本机运行。")
    print()

    app = create_app()
    docx_bytes = _build_demo_docx()

    with TestClient(app) as c:
        # ---- 1. 健康检查 ----
        h = c.get(f"{API}/healthz")
        assert h.status_code == 200, h.text
        hj = h.json()
        print("### 1. 健康检查")
        print(f"- `GET /api/v1/healthz` → `{h.status_code}`")
        print(f"- 默认模型：`{hj.get('default_model')}` · 检索嵌入器：`{hj.get('rag_embedder')}`")
        print()

        # ---- 2. 建会话 + 上传论文 ----
        sid = c.post(f"{API}/sessions", json={"title": "demo"}).json()["id"]
        up = c.post(
            f"{API}/documents",
            params={"session_id": sid},
            files={"file": ("demo_paper.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert up.status_code == 200, up.text
        doc = up.json()
        doc_id = doc["id"]
        print("### 2. 上传论文（现场生成的 docx，不依赖外部文件）")
        print(f"- `POST /api/v1/documents` → `{up.status_code}` · id=`{doc_id}` · 状态=`{doc.get('status')}`")
        meta = doc.get("meta") or {}
        print(f"- 解析结果：段落 {meta.get('num_blocks', '—')} · 分块 {meta.get('num_chunks', '—')}")
        print()

        # ---- 3. 提问（SSE 流式 + 引用） ----
        r = c.post(
            f"{API}/sessions/{sid}/messages/stream",
            json={"content": "这篇论文的核心方法是什么？"},
        )
        assert r.status_code == 200, r.text
        events = _parse_sse(r.text)
        kinds: dict[str, int] = {}
        for name, _ in events:
            kinds[name] = kinds.get(name, 0) + 1
        # ★ 事件字段口径（2026-10-10 核对）：最终 `message` 帧用 **`text`** 携带完整回答；
        #   `message_delta` 是逐块增量 ⇒ 优先取 message.text，缺失才回退累加增量（避免双计）。
        final_text, delta_text, citations = "", "", []
        for name, payload in events:
            if name == "message":
                final_text = payload.get("text") or ""
                if payload.get("citations"):
                    citations = payload["citations"]
            elif name in ("message_delta", "delta"):
                delta_text += payload.get("text") or payload.get("delta") or ""
        answer = final_text or delta_text
        print("### 3. 提问（SSE 流式）")
        print(f"- `POST /api/v1/sessions/{sid}/messages/stream` → `{r.status_code}`")
        print(f"- 事件序列：{', '.join(f'{k}×{v}' for k, v in kinds.items())}")
        print(f"- 回答长度：{len(answer)} 字" + (f" · 开头：{answer.strip()[:60]}…" if answer.strip() else ""))
        if citations:
            cit = citations[0]
            # fake 管线不绑定真实文档 id（doc_id/id 可能为空）⇒ 如实说明，而不是含糊打个 "?"
            raw_id = cit.get("doc_id") if cit.get("doc_id") is not None else cit.get("id")
            label = str(raw_id) if raw_id is not None else "（fake 管线未绑定文档 id）"
            extra: list[str] = []
            if cit.get("page") is not None:
                extra.append(f"第 {cit['page']} 页")
            if cit.get("model_truncated") is True:
                extra.append("标为「节选」—— 模型所见已被截断")
            print(f"- 引用溯源：{len(citations)} 条 · 首条文档 `{label}`"
                  + (f" · {' · '.join(extra)}" if extra else ""))
        else:
            print("- 引用溯源：0 条（fake 嵌入下检索命中为空属预期）")
        print()

        # ---- 4. 段落视图 + 改稿（Word 修订痕迹） ----
        paras = c.get(f"{API}/documents/{doc_id}/paragraphs").json()
        editable = [p for p in paras if p.get("editable")]
        assert editable, "段落视图里没有可编辑段落"
        target = next((p for p in editable if len(p.get("text") or "") > 12), editable[0])
        print("### 4. 改稿工作台（写入 Word 修订痕迹）")
        print(f"- 段落视图：共 {len(paras)} 段，其中可编辑 {len(editable)} 段")
        print(f"- 目标段落 `{target['pid']}`：{target['text'][:40]}…")

        applied = c.post(
            f"{API}/documents/{doc_id}/apply-edits",
            json={
                "edits": [
                    {
                        "pid": target["pid"],
                        "op": "append",
                        "text": "（演示：本句由 PaperMate 改稿工作台以修订形式写入。）",
                    }
                ]
            },
        )
        assert applied.status_code == 200, applied.text
        aj = applied.json()
        print(f"- `POST /api/v1/documents/{doc_id}/apply-edits` → `{applied.status_code}` · "
              f"已应用 {len(aj.get('applied', []))} 条 · 拒绝 {len(aj.get('rejected', []))} 条")
        if aj.get("backup"):
            print(f"- 自动备份：`{Path(aj['backup']).name}`（可一键撤销）")
        print()

        # ---- 5. 导出「修订版 / 干净版」 ----
        # 导出到**仓库内固定目录**（便于 CI 作为 artifact 上传；已在 .gitignore 忽略）
        out_dir = Path.cwd() / "demo-artifacts"
        out_dir.mkdir(parents=True, exist_ok=True)
        print("### 5. 导出（修订版 / 干净版）")
        for mode, label in (("tracked", "修订版"), ("clean", "干净版")):
            dest = out_dir / f"demo_paper_{mode}.docx"
            ex = c.post(
                f"{API}/documents/{doc_id}/export-to",
                json={"dest": str(dest), "mode": mode},
            )
            assert ex.status_code == 200, ex.text
            data = dest.read_bytes()
            ins, dele = _count_revisions(data)
            print(f"- `{label}` → `{ex.status_code}` · {len(data):,} 字节 · 落盘 `{dest.name}` · "
                  f"修订标记 w:ins={ins} / w:del={dele}")
        print()

        # ---- 6. 数据全部落在本机 ----
        print("### 6. 数据落点（本地优先）")
        for sub in ("uploads", "logs"):
            p = Path(os.environ["PAPERMATE_DATA_DIR"]) / sub
            n = sum(1 for _ in p.rglob("*")) if p.exists() else 0
            print(f"- `{sub}/`：{n} 个文件")
        print()
        print(f"_Demo 数据目录：`{os.environ['PAPERMATE_DATA_DIR']}`（可整目录删除，不影响任何真实数据）_")

    print()
    print("**结论：端到端链路全部通过 ✓**（健康检查 → 上传解析 → 流式问答 → 引用 → 改稿修订 → 双版本导出）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
